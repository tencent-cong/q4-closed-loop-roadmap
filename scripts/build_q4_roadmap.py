import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

src = "/Users/congjiang/Downloads/抖音蝉妈妈-日用百货 25 年Q4数据.xlsx"
out = "/Users/congjiang/With/20260923/6fia/闭环链路_Q4节点品类RoadMap.xlsx"
df = pd.read_excel(src)

def classify(c):
    c = str(c)
    if c.startswith("家居家纺-"):
        return "家居家纺"
    if any(x in c for x in ["摆件", "香薰", "线香", "装饰画", "工艺品", "香囊"]):
        return "家居工艺品"
    if any(x in c for x in ["炒锅", "煎锅", "平底锅", "汤锅", "蒸锅", "奶锅", "炖锅", "压力锅", "菜刀", "砧板", "菜板", "剪刀", "瓜果刀", "刨", "锅具"]):
        return "厨具"
    if any(x in c for x in ["餐具", "水具", "保温杯", "保温壶", "马克杯", "玻璃杯", "塑杯", "筷子", "饭盒", "保温桶", "茶具", "酒杯", "碗"]):
        return "餐具水具"
    if any(x in c for x in ["收纳", "衣架", "晾晒", "置物架", "鞋架", "压缩袋", "防尘罩", "挂钩", "粘钩", "储物"]):
        return "收纳用品"
    if any(x in c for x in ["清洁", "拖把", "抹布", "垃圾袋", "垃圾桶", "花洒", "马桶", "扫把", "刷子", "粘毛", "黏毛", "家务手套", "除尘掸", "洗衣膏"]):
        return "清洁工具"
    if any(x in c for x in ["艾灸", "足贴", "眼罩", "护腰", "护膝", "护肩", "护颈", "暖宝宝", "保暖", "热水袋", "按摩", "足浴", "泡脚", "修脚", "护理贴", "暖腹", "暖手贴", "暖脚"]):
        return "功效品"
    return "其他"

df["赛道"] = df["所属类目"].map(classify)
df["细分类目"] = df["所属类目"].astype(str).str.split("-").str[-1]
tracks = ["功效品", "清洁工具", "收纳用品", "其他", "厨具", "餐具水具", "家居家纺", "家居工艺品"]
top = {}
for t in tracks:
    g = df[df["赛道"] == t].groupby("细分类目").agg(SKU=("商品", "size"), 销售额指数=("销售额-指数", "sum")).sort_values("销售额指数", ascending=False).head(8)
    top[t] = "、".join(g.index.tolist())

road = {
"功效品": [
"蒸汽眼罩、足贴、护腰护膝、暖宝宝", 
"蒸汽眼罩、护腰护膝、热敷贴，切换秋冬保暖心智", 
"暖宝宝/暖贴囤货、足贴、蒸汽眼罩、肩颈腰膝贴", 
"足贴、泡脚/足浴、热水袋、发热护具", 
"暖贴、足浴桶/泡脚包、热水袋、护腰护膝，兼顾长辈礼赠", 
"ADQ已投36个SKU、累计消耗902.3万；已验证艾灸贴/足贴/蒸汽眼罩等。外部待测八髎贴、三阴交贴、乳贴、鼻安合香、颈椎贴/肩周贴及秋冬骨贴、膝腰贴、老黑膏、透骨膏等", 
"基本盘+季节增量；P0"
],
"清洁工具": [
"拖把、花洒、水龙头过滤器、家务手套", 
"浴室/地面清洁、玻璃结露清洁、羽绒服局部护理", 
"拖把、抹布、垃圾袋、清洁刷、花洒套装囤货", 
"地暖地面清洁、门窗冷凝水工具、火锅后去油工具", 
"年终大扫除：拖把、扫把、缝隙刷、玻璃擦、油污清洁", 
"闭环基本盘集中拖把；非闭环水龙头过滤器、花洒较强。闭环可补厨房刷、缝隙刷、除尘掸等低客单强演示品", 
"基本盘扩品；P0"
],
"收纳用品": [
"压缩袋、收纳箱、厨房置物架、衣物收纳", 
"夏物/薄被换季收纳：压缩袋、被褥袋、防尘罩", 
"压缩袋、收纳箱、厨房置物架、鞋架组合装", 
"厚衣/被褥收纳、衣柜分区、桌面整理", 
"年终整理：收纳箱、保鲜盒、厨房置物架、鞋架", 
"闭环压缩袋、厨房置物架已有消耗；非闭环保鲜盒、厨房置物架、收纳箱/收纳包持续强，适合向闭环迁移", 
"换季+大促双峰；P0"
],
"其他": [
"居家鞋、口罩、一次性手套、浴巾、居家日用套装", 
"棉拖/居家鞋、口罩、浴巾、保暖手套", 
"居家鞋、一次性手套、浴巾、修毛器/鼻毛修剪器", 
"棉拖、暖脚用品、修毛器、鞋靴护理工具", 
"防寒居家鞋、浴巾礼盒、个人护理工具礼盒", 
"Q4表内居家鞋、口罩、一次性手套规模靠前；修毛器适合用毛衣起球/年底焕新素材补充测试", 
"工具型机会；P1"
],
"厨具": [
"炒锅、菜刀、砧板、平底锅", 
"换新厨房：炒锅、菜刀、砧板套装", 
"炒锅、平底锅、菜刀、砧板、蒸锅大促换新", 
"火锅/炖煮：汤锅、蒸锅、炖锅、压力锅", 
"年末聚餐：炒锅、蒸锅、炖锅、刀具套装", 
"年度热力图显示炒锅全年强基本盘；菜刀在11月明显走强，砧板/菜板、平底锅为稳定次核心", 
"常青大盘+聚餐增量；P0"
],
"餐具水具": [
"保温杯、保温壶、筷子、玻璃杯", 
"保温杯/壶、焖烧杯、饭盒", 
"保温杯礼盒、餐具套装、筷子、玻璃杯", 
"保温杯、保温壶、焖烧杯、火锅餐具", 
"冬至热饮：保温杯/壶；年末礼赠餐具水具套装", 
"Q4表内保温杯为绝对头部，筷子、保温壶为第二梯队；适合保温功能+礼赠双素材", 
"冬季刚需+礼赠；P0"
],
"家居家纺": [
"床品套件、棉花被、化纤被、枕类", 
"厚被/冬被、床品套件、毛毯、床垫", 
"冬被、床品四件套、枕芯、毛毯大促换新", 
"羽绒被/羊毛被、绒毯、保暖床垫、坐垫", 
"深冬加厚：冬被、毛毯、床品套件；年末家纺礼盒", 
"年度热力图显示床单/床笠/隔脏睡袋和枕头为全年基本盘；被芯、羽绒被、纤维被在秋冬及大促放大", 
"Q4最大确定性；P0"
],
"家居工艺品": [
"香薰、线香、香薰蜡烛、香囊", 
"秋日香氛、乔迁摆件、居家氛围", 
"香薰礼盒、香器、装饰摆件", 
"冬季居家氛围：香薰蜡烛、线香、香器", 
"跨年/乔迁/送礼：香薰礼盒、吉祥摆件、香囊", 
"Q4表内车载香薰规模最高，线香、香薰蜡烛为居家机会；节点内容只讲感官、氛围和礼赠，避免功效及超自然宣称", 
"礼赠+氛围增量；P1"
]
}

wb = Workbook()
ws = wb.active
ws.title = "Q4 Road Map"
headers = ["赛道", "Q4大盘SKU", "销售额指数占比", "Q4大盘头部细分品", "10月换季/国庆", "双11", "立冬/小雪", "双12·大雪/冬至/年末", "外部跑品与闭环判断", "闭环角色/优先级"]
ws.append(["闭环链路 Q4 节点品类 Road Map"])
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
ws.append(["数据底座：抖音蝉妈妈2025年Q4商品TOP1500；外部证据：厨具/家纺年度热力图、清洁/收纳闭环与非闭环热力图、功效品投放矩阵。赛道口径参考情报台创意与卖点分析。"])
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
ws.append(headers)
for i, t in enumerate(tracks, start=4):
    vals = road[t]
    ws.append([t, f'=COUNTIF(商品明细!Q:Q,A{i})', f'=SUMIF(商品明细!Q:Q,A{i},商品明细!M:M)/SUM(商品明细!M:M)', top[t], vals[1], vals[2], vals[3], vals[4], vals[5], vals[6]])

ws2 = wb.create_sheet("大盘数据")
ws2.append(["赛道", "SKU数", "销售额指数", "销售额指数占比", "头部细分类目"])
for i, t in enumerate(tracks, start=2):
    ws2.append([t, f'=COUNTIF(商品明细!Q:Q,A{i})', f'=SUMIF(商品明细!Q:Q,A{i},商品明细!M:M)', f'=C{i}/SUM(C$2:C$9)', top[t]])
ws2.append(["合计", "=SUM(B2:B9)", "=SUM(C2:C9)", "=SUM(D2:D9)", ""])

ws3 = wb.create_sheet("商品明细")
cols = list(df.columns)
ws3.append(cols)
for row in df.itertuples(index=False, name=None):
    ws3.append(list(row))

ws4 = wb.create_sheet("节点总览")
ws4.append(["阶段", "建议时间", "节点目标", "主攻赛道", "代表细分品"])
node_rows = [
["换季启动", "9月下旬-10月中旬", "承接降温、换季整理和居家焕新", "家居家纺/收纳/功效品/其他", "冬被、床品套件、压缩袋、暖贴、居家鞋"],
["双11蓄水与爆发", "10月中旬-11月11日", "全年最大规模节点，基本盘囤货+秋冬新品放量", "全赛道", "冬被、炒锅刀具、拖把清洁套装、压缩袋、暖贴、保温杯"],
["立冬/小雪", "11月7日-11月22日", "从换季切入严寒刚需与热饮热食场景", "功效品/家纺/餐具水具/厨具", "暖宝宝、足贴、护腰护膝、冬被、保温杯、汤锅蒸锅"],
["双12", "11月下旬-12月12日", "承接双11补货和礼赠需求", "家纺/收纳/清洁/家居工艺品", "家纺套件、收纳组合、清洁囤货、香薰礼盒"],
["大雪/冬至/年末", "12月7日-12月31日", "深冬保暖、年终整理、聚餐及跨年礼赠", "全赛道", "深冬保暖品、年终大扫除、锅具餐具、香薰摆件、个人护理工具"],
["春节前置", "12月下旬起", "提前蓄水年货、焕新、清洁和长辈礼赠", "清洁/厨具/餐具水具/功效品/家居工艺品", "大扫除工具、锅具套装、餐具礼盒、足浴/热敷、吉祥摆件"]]
for r in node_rows:
    ws4.append(r)

ws5 = wb.create_sheet("口径说明")
notes = [
["项目", "说明"],
["赛道", "功效品、清洁工具、收纳用品、生活日用其他、厨具、餐具水具、家居家纺、家居工艺品，共8个赛道。"],
["Q4大盘", "读取蝉妈妈2025年Q4商品TOP1500。SKU数为表内商品数；销售额指数占比为各赛道销售额指数之和/全表销售额指数之和。"],
["映射", "因蝉妈妈类目名称与情报台最新统一类目存在版本差异，本稿按商品类目关键词近似映射；建议最终上线前抽样复核。"],
["外部跑品", "来自用户提供的5张图，只作为方向性证据。热力图颜色代表相对消耗，不等同于绝对市场规模；图中未投品不代表已验证。"],
["功效品合规", "节点文案优先描述材质、热敷、保暖、使用场景和感官体验；避免治疗、根治、助眠及其他无依据功效。"],
["闭环节奏", "先用小店/直播验证点击率、转化率与退货售后，再复制到更大闭环预算；P0先建计划，P1作为补充测试。"]]
for r in notes:
    ws5.append(r)

navy = "17365D"
blue = "DCEBFA"
orange = "FCE8D5"
green = "DDEFE2"
white = "FFFFFF"
thin = Side(style="thin", color="D8DEE8")
for sh in wb.worksheets:
    sh.sheet_view.showGridLines = False
    sh.freeze_panes = "A4" if sh.title == "Q4 Road Map" else "A2"
    for row in sh.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="1F2937")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    header_row = 3 if sh.title == "Q4 Road Map" else 1
    for cell in sh[header_row]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(name="Arial", bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sh.auto_filter.ref = sh.dimensions

ws[1][0].fill = PatternFill("solid", fgColor=navy)
ws[1][0].font = Font(name="Arial", size=18, bold=True, color=white)
ws[1][0].alignment = Alignment(horizontal="left", vertical="center")
ws[2][0].fill = PatternFill("solid", fgColor="EAF1F8")
ws[2][0].font = Font(name="Arial", size=10, italic=True, color="40546D")
ws.row_dimensions[1].height = 34
ws.row_dimensions[2].height = 36
ws.row_dimensions[3].height = 35
for r in range(4, 12):
    ws.row_dimensions[r].height = 95
    ws.cell(r, 1).fill = PatternFill("solid", fgColor=green if ws.cell(r,1).value in ["功效品","清洁工具","收纳用品"] else blue)
    ws.cell(r, 1).font = Font(name="Arial", bold=True, color=navy)
for c, w in enumerate([14, 11, 13, 31, 33, 36, 34, 40, 60, 22], start=1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws.column_dimensions["C"].width = 14
for r in range(4, 12):
    ws.cell(r, 3).number_format = "0.0%"

for c, w in enumerate([16, 12, 18, 18, 65], start=1):
    ws2.column_dimensions[get_column_letter(c)].width = w
for r in range(2, 11):
    ws2.cell(r, 4).number_format = "0.0%"
ws2.conditional_formatting.add("D2:D9", ColorScaleRule(start_type="min", start_color="FFFFFF", mid_type="percentile", mid_value=50, mid_color="DCEBFA", end_type="max", end_color="4F7FEF"))

for c in range(1, ws3.max_column + 1):
    ws3.column_dimensions[get_column_letter(c)].width = 16
ws3.column_dimensions["B"].width = 52
ws3.column_dimensions["E"].width = 48
ws3.column_dimensions["Q"].width = 15
ws3.column_dimensions["R"].width = 24
ws3.freeze_panes = "A2"

for c, w in enumerate([18, 23, 38, 38, 70], start=1):
    ws4.column_dimensions[get_column_letter(c)].width = w
for r in range(2, ws4.max_row + 1):
    ws4.row_dimensions[r].height = 48
for c, w in enumerate([20, 110], start=1):
    ws5.column_dimensions[get_column_letter(c)].width = w
for r in range(2, ws5.max_row + 1):
    ws5.row_dimensions[r].height = 45

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.save(out)
print(out)
